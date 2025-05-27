import docx
import logging
import os
from typing import Tuple, Dict, Optional # For type hinting
from docx.document import Document as DocxDocument
from docx.table import Table, _Cell as CellType # For type hinting Table and Cell
from docx.text.paragraph import Paragraph # For type hinting Paragraph
from docx.image.exceptions import UnrecognizedImageError
# from docx.text.run import Run # Not directly used in the refactored main loop
# from docx.shape import InlineShape # Not directly used in the refactored main loop

logger = logging.getLogger(__name__)

# Namespace dictionary for XPath queries
ns = {
    'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
    'wp': 'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing',
    'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
    'pic': 'http://schemas.openxmlformats.org/drawingml/2006/picture',
    'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
}

def _process_paragraph_content(para: Paragraph, doc_part, images_sub_dir: str, image_counter: int, extracted_data: Dict[str, str]) -> Tuple[str, int]:
    """Helper function to process a paragraph's content for text and images."""
    paragraph_text_parts = []
    for run in para.runs:
        if run.text:
            paragraph_text_parts.append(run.text)
        
        # Image extraction logic (similar to before, adapted for this helper)
        for drawing in run.element.findall('.//w:drawing', namespaces=ns):
            inline_elements = drawing.findall('.//wp:inline', namespaces=ns)
            for inline in inline_elements:
                graphic = inline.find('.//a:graphic', namespaces=ns)
                if graphic is not None:
                    graphic_data = graphic.find('.//a:graphicData', namespaces=ns)
                    if graphic_data is not None:
                        pic = graphic_data.find('.//pic:pic', namespaces=ns)
                        if pic is not None:
                            blip_fill = pic.find('.//pic:blipFill', namespaces=ns)
                            if blip_fill is not None:
                                blip = blip_fill.find('.//a:blip', namespaces=ns)
                                if blip is not None:
                                    rId = blip.get(f'{{{ns["r"]}}}embed')
                                    if rId and doc_part: # Ensure doc_part is available
                                        try:
                                            image_part = doc_part.related_parts[rId]
                                            image_bytes = image_part.blob
                                            image_ext = image_part.default_image_ext
                                            
                                            image_counter += 1
                                            image_filename = f"image_{image_counter}.{image_ext}"
                                            image_save_path = os.path.join(images_sub_dir, image_filename)
                                            
                                            with open(image_save_path, 'wb') as img_file:
                                                img_file.write(image_bytes)
                                            logger.info(f"成功提取并保存图片: {image_save_path}")
                                            
                                            placeholder = f"[IMAGE_PLACEHOLDER:{image_filename}]"
                                            paragraph_text_parts.append(placeholder)
                                            # Store image info (filename for now)
                                            extracted_data[placeholder] = image_filename 
                                        except KeyError:
                                            logger.warning(f"在文档部件中未找到图片 rId: {rId}")
                                        except IOError as e:
                                            logger.error(f"无法写入图片文件 {image_save_path}: {e}")
                                        except Exception as e:
                                            logger.error(f"保存图片 {image_save_path} 时发生未知错误: {e}")
    return "".join(paragraph_text_parts), image_counter

def _table_to_markdown(table: Table) -> str:
    """Converts a docx.table.Table object to a Markdown string."""
    markdown_lines = []
    if not table.rows:
        return ""
        
    # Process header row
    header_cells = table.rows[0].cells
    header_texts = [cell.text.strip().replace('\n', ' ').replace('|', r'\|') for cell in header_cells]
    markdown_lines.append("| " + " | ".join(header_texts) + " |")
    markdown_lines.append("| " + " | ".join(["---"] * len(header_cells)) + " |")
    
    # Process data rows
    for row in table.rows[1:]:
        data_cells = row.cells
        data_texts = [cell.text.strip().replace('\n', ' ').replace('|', r'\|') for cell in data_cells]
        markdown_lines.append("| " + " | ".join(data_texts) + " |")
        
    return "\n".join(markdown_lines)

def extract_content_from_docx(file_path: str, results_dir: str) -> Optional[Tuple[str, Dict[str, str]]]:
    """
    从 .docx 文件中提取文本、图片和表格。
    图片被保存，文本中的图片和表格被占位符替换。

    返回:
        一个元组 (text_with_placeholders, extracted_data)，其中:
        - text_with_placeholders: 包含所有占位符的文本字符串。
        - extracted_data: 一个字典，将占位符映射到其实际内容
                          (图片是文件名，表格是其 Markdown 字符串)。
        如果发生错误，则返回 None。
    """
    try:
        logger.info(f"开始从 DOCX 文件提取内容 (包括图片和表格): {file_path}")
        doc: DocxDocument = docx.Document(file_path)
        
        overall_content_parts = [] # Stores text, image placeholders, and table placeholders in order
        extracted_data: Dict[str, str] = {} 
        image_counter = 0
        table_counter = 0
        
        images_sub_dir = os.path.join(results_dir, 'images')
        if not os.path.exists(images_sub_dir):
            os.makedirs(images_sub_dir, exist_ok=True)
            logger.info(f"创建图片子目录: {images_sub_dir}")

        # Iterate through block-level items (paragraphs and tables)
        # doc.element.body contains the main document content.
        # Children of body can be paragraphs (<w:p>) or tables (<w:tbl>).
        for child_element in doc.element.body:
            if child_element.tag == f'{{{ns["w"]}}}p': # Paragraph
                para = Paragraph(child_element, doc)
                para_text, image_counter = _process_paragraph_content(
                    para, doc.part, images_sub_dir, image_counter, extracted_data
                )
                if para_text: # Only add if paragraph has actual content (text or image placeholder)
                    overall_content_parts.append(para_text)
            
            elif child_element.tag == f'{{{ns["w"]}}}tbl': # Table
                table = Table(child_element, doc)
                table_counter += 1
                table_placeholder = f"[TABLE_PLACEHOLDER_{table_counter}]"
                
                markdown_table_str = _table_to_markdown(table)
                if markdown_table_str: # Only add if table conversion yields something
                    extracted_data[table_placeholder] = markdown_table_str
                    overall_content_parts.append(table_placeholder) 
                else:
                    logger.warning(f"表格 {table_counter} 转换为空 Markdown，已跳过。")


        final_text_with_placeholders = "\n\n".join(overall_content_parts)
        
        if image_counter == 0:
            logger.info(f"DOCX 文件 {file_path} 中未找到或提取图片。")
        if table_counter == 0:
            logger.info(f"DOCX 文件 {file_path} 中未找到或提取表格。")
        
        logger.info(f"成功从 DOCX 文件提取内容: {image_counter} 张图片, {table_counter} 个表格。文件: {file_path}")
        return final_text_with_placeholders, extracted_data

    except docx.opc.exceptions.PackageNotFoundError:
        logger.error(f"无法打开或解析 DOCX 文件 (文件不存在、已损坏或格式无效): {file_path}")
        return None
    except UnrecognizedImageError as uie:
        logger.error(f"从 DOCX 文件 {file_path} 提取图片时遇到无法识别的图片类型: {uie}")
        return None
    except Exception as e:
        logger.error(f"从 DOCX 文件 {file_path} 提取内容时发生意外错误: {e}", exc_info=logger.isEnabledFor(logging.DEBUG))
        return None